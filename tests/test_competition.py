"""Offline acquisition checks; small XLSX fixtures are synthetic table tests only."""
import copy
import csv
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

from competition.kba import count, parse
from competition import source_path
from competition.derive import ROOT, derive, matches, positive, ratio, validate, verify_raw


class KbaCompetitionTests(unittest.TestCase):
    def workbook(self, folder, legend=True, mutate=None):
        book = openpyxl.Workbook()
        table = book.active
        table.title = 'FZ 10.1'
        header = [None]*39
        header[1:4] = ['Marke','Modellreihe','Insgesamt']
        header[37] = 'mit Elektroantrieb (BEV)'
        table.append(header)
        labels = [None]*39
        labels[3:5] = ['November 2024','Jan. - Nov. 2024']
        labels[37:39] = labels[3:5]
        table.append(labels)
        for brand, model, month, ytd, electric, electric_ytd in [
            ('VW','ID.3',10,100,6,60),('VW ZUSAMMEN',None,10,100,6,60),
            ('MORGAN',None,5,50,'-','-'),('NEUZULASSUNGEN INSGESAMT',None,15,150,6,60)]:
            row = [None]*39
            row[1:5] = [brand,model,month,ytd]
            row[37:39] = [electric,electric_ytd]
            table.append(row)
        notes = book.create_sheet('Impressum')
        if legend:
            notes.append(['-','nichts vorhanden oder keine Veränderung'])
        notes.append(['Datenlizenz Deutschland – Namensnennung – Version 2.0'])
        if mutate:
            mutate(book)
        path = Path(folder)/'table.xlsx'
        book.save(path)
        record = {'id':'test-table','period':'2024-11','sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
        return path, record

    def test_counts_reject_invalid_and_do_not_fill_suppression(self):
        for value in [-1, 1.5, True, '1', float('inf'), float('nan'), '=SUM(A1)']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                count(value, True)
        for value in [None,'','/','X','...','-']:
            self.assertIsNone(count(value))
        self.assertEqual(count('-',True),0)

    def test_standalone_brand_reconciles_month_and_ytd(self):
        with tempfile.TemporaryDirectory() as tmp:
            p,r = self.workbook(tmp)
            result = parse(p,r)
        self.assertEqual(len(result['brands']),4)
        self.assertEqual({x['scope'] for x in result['checks']['brand_reconciliation']},{'monthly','ytd'})
        self.assertTrue(all(x['status']=='passed' for x in result['checks']['brand_reconciliation']))

    def test_missing_legend_preserves_dash_unknown(self):
        with tempfile.TemporaryDirectory() as tmp:
            p,r = self.workbook(tmp,legend=False)
            result = parse(p,r)
        m = next(x for x in result['brands'] if x['brand']=='MORGAN' and x['powertrain']=='BEV')
        self.assertIsNone(m['monthly'])
        self.assertIsNone(result['checks']['dash_zero_evidence'])

    def test_bad_hash_and_wrong_period_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            p,r = self.workbook(tmp)
            with self.assertRaisesRegex(ValueError,'hash'):
                parse(p,{**r,'sha256':'0'*64})
            with self.assertRaisesRegex(ValueError,'period mismatch'):
                parse(p,{**r,'period':'2024-10'})

    def test_wrong_header_and_wrong_totals_rejected(self):
        for mutate in [lambda b:setattr(b.active['AL1'],'value','BEV unknown'),
                       lambda b:setattr(b.active['E6'],'value',149)]:
            with self.subTest(mutate=mutate), tempfile.TemporaryDirectory() as tmp:
                p,r = self.workbook(tmp,mutate=mutate)
                with self.assertRaises(ValueError):
                    parse(p,r)

    def test_january_standalone_monthly_total_not_lost(self):
        def january(book):
            table = book.active
            for c in ['D2','AL2']:
                table[c] = 'Januar 2024'
            for row in range(3,7):
                table.cell(row,5).value = None
                table.cell(row,39).value = None
        with tempfile.TemporaryDirectory() as tmp:
            p,r = self.workbook(tmp,mutate=january)
            result = parse(p,{**r,'period':'2024-01'})
        self.assertEqual(len(result['brands']),4)
        self.assertTrue(any(x['status']=='total_not_published_no_imputation' for x in result['checks']['brand_reconciliation']))


class PriceResearchTests(unittest.TestCase):
    def setUp(self):
        self.offers = json.loads((ROOT/'price_observations.json').read_text())['observations']
        self.sources = json.loads((ROOT/'price_sources.json').read_text())['sources']

    def test_actual_catalog_and_source_bindings_validate(self):
        validate(self.offers,self.sources)

    def test_invalid_numbers_dates_and_sources_rejected(self):
        for field,value in [('quoted_price_eur','NaN'),('quoted_price_eur','-1'),
                            ('range_km_upper_or_bound',0),('price_date','2024-13-01'),
                            ('source_id','not-a-source'),('range_km_lower',99999)]:
            row = copy.deepcopy(next(r for r in self.offers if r['range_km_upper_or_bound']))
            row[field] = value
            with self.subTest(field=field),self.assertRaises(ValueError):
                validate([row],self.sources)
        with self.assertRaises(ValueError):
            validate([self.offers[0],self.offers[0]],self.sources)
        for value in [float('inf'),float('nan'),True,0]:
            with self.assertRaises(ValueError):
                positive(value)

    def test_price_range_arithmetic_and_restrictions(self):
        row = copy.deepcopy(next(r for r in self.offers if r['source_id']=='bmw_ix1_2023'))
        self.assertEqual(ratio(row)['eur_per_100km_declared_max_range'],'10084.21')
        self.assertEqual(ratio(row)['eur_per_100km_declared_min_range'],'11139.53')
        for field,value in [('range_cycle','NEDC'),('battery_in_price','rented_excluded'),
                            ('price_kind','promotion_after_manufacturer'),('range_operator','gt'),
                            ('vat','not_stated'),('quoted_price_operator','lt')]:
            changed = {**row,field:value}
            with self.subTest(field=field):
                self.assertEqual(ratio(changed)['eur_per_100km_declared_max_range'],'')

    def test_no_historical_price_forward_fill(self):
        row = next(r for r in self.offers if r['source_id']=='bmw_ix1_2023')
        model = {'year':'2023','kba_brand':'BMW','model_family':'X1'}
        self.assertTrue(matches(row,model))
        self.assertFalse(matches(row,{**model,'year':'2025'}))
        self.assertFalse(matches(row,{**model,'kba_brand':'VW'}))

    def test_highest_price_not_automatically_longest_range(self):
        rows = [r for r in self.offers if r['source_id']=='cupra_tavascan_2024']
        entry = next(r for r in rows if r['role_in_source_snapshot']=='entry')
        top = next(r for r in rows if r['role_in_source_snapshot']=='highest_price_in_report')
        self.assertGreater(Decimal(top['quoted_price_eur']),Decimal(entry['quoted_price_eur']))
        self.assertLess(top['range_km_upper_or_bound'],entry['range_km_upper_or_bound'])

    def test_derived_views_are_deterministic_with_explicit_gaps(self):
        with tempfile.TemporaryDirectory() as a,tempfile.TemporaryDirectory() as b:
            report = derive(ROOT,a)
            self.assertFalse(report['full_year_price_panel_complete'])
            self.assertEqual(report,derive(ROOT,b))
            for p in Path(a).iterdir():
                self.assertEqual(p.read_bytes(),(Path(b)/p.name).read_bytes())
            with (Path(a)/'price_coverage_by_model_year.csv').open() as handle:
                rows = list(csv.DictReader(handle))
        self.assertEqual(len(rows),80)
        self.assertTrue(any(r['coverage_status']=='price_missing_no_backfill' for r in rows))
        self.assertTrue(all(r['full_year_entry_and_top_coverage']=='False' for r in rows))

    def test_raw_integrity_check_fails_on_manipulation(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            (p/'raw').write_bytes(b'actual')
            record = {'id':'x','file':'raw','bytes':6,'sha256':hashlib.sha256(b'actual').hexdigest()}
            (p/'sources.json').write_text(json.dumps({'sources':[record]}))
            self.assertEqual(verify_raw([p/'sources.json'],p),1)
            (p/'raw').write_bytes(b'edited')
            with self.assertRaises(ValueError):
                verify_raw([p/'sources.json'],p)

    def test_source_paths_cannot_escape_raw_cache(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)/'raw'
            root.mkdir()
            outside = Path(tmp)/'private.txt'
            outside.write_text('not source data')
            (root/'link').symlink_to(outside)
            for name in ['../private.txt','/tmp/private.txt','..\\private.txt',
                         '%2e%2e%2fprivate.txt','link']:
                with self.subTest(name=name),self.assertRaises(ValueError):
                    source_path(root,name)
