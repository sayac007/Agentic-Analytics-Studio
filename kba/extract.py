"""Versioned FZ10 layouts. Cell coordinates are 1-based; never guess columns."""

import datetime as dt
import hashlib
import io
import math
import re

MONTHS = (
    "Januar Februar März April Mai Juni Juli August September Oktober November Dezember"
).split()
TITLE = "Neuzulassungen von Personenkraftwagen nach Marken und Modellreihen"
LICENSE_SHA = "7d2235cbbf83fe3107dcd678e39831abd60f22212c0587ec3989f99aa8cdbabe"
PROFILES = {
    "fz10-xls-2018-v1": dict(
        format="xls",
        sheets=[
            "Pkw-NZL Marken, Modellreihen",
            "FZ10.1  ",
            "FZ10.1 ",
            "FZ 10.1  ",
            "FZ 10.1",
        ],
        header=8,
        bev=None,
        first="2018-01",
        last="2018-11",
    ),
    "fz10-xlsx-legacy-v1": dict(
        format="xlsx",
        sheets=[
            "FZ 10.1",
            "FZ 10.1  ",
            "FZ10.1",
            "FZ10.1 ",
            "FZ10.1  ",
            "FZ 10.1 ",
            "FZ 10-November 2020  ",
        ],
        header=8,
        bev=None,
        first="2018-12",
        last="2020-12",
    ),
    "fz10-xlsx-bev-v1": dict(
        format="xlsx",
        sheets=["FZ10.1", "FZ 10.1"],
        header=8,
        bev=37,
        first="2021-01",
        last="2025-12",
    ),
    "fz10-xlsx-bev-202408-v1": dict(
        format="xlsx",
        sheets=["FZ 10.1"],
        header=9,
        bev=37,
        first="2024-08",
        last="2024-08",
    ),
}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def clean(value):
    return " ".join(str(value or "").split())


def month(value):
    if (
        not isinstance(value, str)
        or re.fullmatch(r"[0-9]{4}-(0[1-9]|1[0-2])", value) is None
    ):
        raise ValueError("Invalid month")
    dt.date.fromisoformat(value + "-01")
    return value


def periods(start, end):
    month(start)
    month(end)
    if start > end:
        raise ValueError("Reversed range")
    y, m = map(int, start.split("-"))
    result = []
    while f"{y:04}-{m:02}" <= end:
        result.append(f"{y:04}-{m:02}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return result


def count(value):
    # Even the official '-' symbol is not silently converted to zero.
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
        or value < 0
        or value != int(value)
    ):
        raise ValueError(f"Not an explicit nonnegative integer count: {value!r}")
    return int(value)


def publication(value):
    text = clean(value)
    for i, name in enumerate(MONTHS, 1):
        m = re.fullmatch(rf"(?:(\d{{1,2}})\. )?{name} (\d{{4}})", text)
        if m:
            year = int(m[2])
            day = int(m[1]) if m[1] else None
            if day is not None:
                return {
                    "value": dt.date(year, i, day).isoformat(),
                    "precision": "day",
                    "raw": text,
                }
            return {"value": f"{year:04}-{i:02}", "precision": "month", "raw": text}
    if not text:
        return {"value": None, "precision": "unknown", "raw": None}
    raise ValueError("Unrecognized publication date: " + text)


def parse(raw, record):
    """Validate raw hash, explicit profile, month/title/header, then exact total rows."""
    if sha(raw) != record["sha256"]:
        raise ValueError("Raw hash mismatch")
    profile = PROFILES[record["parser_profile"]]
    period = month(record["period"])
    if not profile["first"] <= period <= profile["last"] or (
        period == "2024-08" and profile["header"] != 9
    ):
        raise ValueError("Profile outside reviewed period")
    if profile["format"] == "xlsx":
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
        names = book.sheetnames

        def get(sheet, row, col):
            c = book[sheet].cell(row, col)
            if c.data_type == "f":
                raise ValueError("Formula in extraction cell")
            return c.value

        def nrows(sheet):
            return book[sheet].max_row
    else:
        import xlrd

        book = xlrd.open_workbook(file_contents=raw)
        names = book.sheet_names()

        def get(sheet, row, col):
            s = book.sheet_by_name(sheet)
            return s.cell_value(row - 1, col - 1) if col <= s.ncols else None

        def nrows(sheet):
            return book.sheet_by_name(sheet).nrows

    matches = [n for n in names if n in profile["sheets"]]
    if len(matches) != 1:
        raise ValueError("Unknown or ambiguous table sheet")
    sheet = matches[0]
    h = profile["header"]
    label = MONTHS[int(period[5:]) - 1] + " " + period[:4]
    if clean(get(sheet, 5, 2)) != (
        TITLE + " im " + label if period in ["2020-11", "2020-12"] else TITLE
    ) or clean(get(sheet, 6, 2)) not in [
        "FZ10.1 " + TITLE + " im " + label,
        "FZ 10.1 " + TITLE + " im " + label,
        *(([TITLE + " im " + label]) if period == "2019-08" else []),
    ]:
        raise ValueError("Unknown table title or reporting period")
    if (
        clean(get(sheet, h, 4)) != "Insgesamt"
        or clean(get(sheet, h + 1, 2))
        != ("Marken" if period in ["2018-06", "2018-07"] else "Marke")
        or clean(get(sheet, h + 1, 3)) not in ["Modellreihe", "Marke/Modellreihe"]
    ):
        raise ValueError("Unknown total/brand/model header layout")
    if (
        clean(get(sheet, h + 1, 4)) != label
        or clean(get(sheet, h, 7)) != "mit Dieselantrieb"
    ):
        raise ValueError("Wrong monthly column")
    if profile["bev"]:
        if (
            clean(get(sheet, h, 37)) != "mit Elektroantrieb (BEV)"
            or clean(get(sheet, h + 1, 37)) != label
        ):
            raise ValueError("BEV definition or month mismatch")
    elif clean(get(sheet, h, 13)) != "mit Elektroantrieb":
        raise ValueError("Unknown legacy drivetrain layout")
    pub = publication(None)
    revision = None
    exclusions = []
    license_cell = None
    if period >= "2020-01" and "Impressum" not in names:
        raise ValueError("Required imprint missing")
    if "Impressum" in names:
        shift = 1 if period == "2025-06" else 0
        report_date = get("Impressum", 9, 3)
        if period == "2022-05" and isinstance(report_date, dt.datetime):
            report_date = MONTHS[report_date.month - 1] + " " + str(report_date.year)
        if (
            clean(report_date) != label
            or clean(get("Impressum", 15 + shift, 3)) != "Kraftfahrt-Bundesamt"
            or clean(get("Impressum", 13 + shift, 3)) != "monatlich"
        ):
            raise ValueError("Imprint mismatch")
        pub_value = get("Impressum", 11, 3)
        if period == "2022-05" and isinstance(pub_value, dt.datetime):
            if book["Impressum"]["C11"].number_format != "mmmm\\ yyyy":
                raise ValueError("Unreviewed date display format")
            pub_value = MONTHS[pub_value.month - 1] + " " + str(pub_value.year)
        pub = publication(pub_value)
        if shift:
            if clean(get("Impressum", 12, 2)) != "Korrigierte Version vom:":
                raise ValueError("Unknown revision layout")
            revision = publication(get("Impressum", 12, 3))
        for r in (26 + shift, 28 + shift):
            if (
                clean(get("Impressum", r, 3))
                == "Datenlizenz Deutschland - Namensnennung - Version 2.0"
            ):
                license_cell = "C" + str(r)
        if license_cell is None:
            raise ValueError("Unknown imprint license")
    observations = []
    seen = set()
    total_rows = []
    for r in range(h + 2, nrows(sheet) + 1):
        name = clean(get(sheet, r, 2))
        model = clean(get(sheet, r, 3))
        if name == "NEUZULASSUNGEN INSGESAMT":
            if model:
                raise ValueError("Market total has model label")
            total_rows.append(r)
        if name.endswith(" ZUSAMMEN") and name != "SONSTIGE ZUSAMMEN":
            if model:
                if (period, r, name, model) == (
                    "2020-08",
                    117,
                    "HONDA ZUSAMMEN",
                    "H-1 STAREX",
                ):
                    exclusions.append(
                        dict(
                            cell="D117",
                            reason="Ambiguous Honda total has another model label in source; excluded",
                        )
                    )
                    continue
                raise ValueError("Brand total has model label")
            brand = name.removesuffix(" ZUSAMMEN")
            if brand in seen:
                raise ValueError("Duplicate brand total")
            seen.add(brand)
            # Suppressed or symbolic counts are not observations; retain an explicit gap.
            value = get(sheet, r, 4)
            if isinstance(value, str) and value.strip() in ["-", ".", "/", "X", ""]:
                exclusions.append(
                    dict(
                        cell="D" + str(r),
                        series="brand:" + brand,
                        reason="Published symbolic value, not converted to zero: "
                        + repr(value),
                    )
                )
                continue
            observations.append(
                dict(
                    series="brand:" + brand,
                    value=count(value),
                    powertrain="all",
                    cell="D" + str(r),
                )
            )
    if len(total_rows) != 1 or "TESLA" not in seen:
        raise ValueError("Missing/duplicate market or Tesla total")
    if profile["bev"]:
        r = total_rows[0]
        observations.append(
            dict(
                series="market:BEV",
                value=count(get(sheet, r, 37)),
                powertrain="BEV_only",
                cell="AK" + str(r),
            )
        )
    if period.endswith("-12"):
        for observation in observations:
            col = 38 if observation["powertrain"] == "BEV_only" else 5
            if "Dezember" not in clean(get(sheet, h + 1, col)) or not clean(
                get(sheet, h + 1, col)
            ).endswith(period[:4]):
                raise ValueError("Unknown December cumulative column")
            row = int(re.search(r"[0-9]+$", observation["cell"])[0])
            value = get(sheet, row, col)
            if isinstance(value, str) and value.strip() in ["-", ".", "/", "X", ""]:
                observation["cumulative_value"] = None
                exclusions.append(
                    dict(
                        cell=("AL" if col == 38 else "E") + str(row),
                        reason="Symbolic annual cumulative unavailable; monthly value retained",
                    )
                )
            else:
                observation["cumulative_value"] = count(value)
            observation["cumulative_cell"] = ("AL" if col == 38 else "E") + str(row)
    return dict(
        sheet=sheet,
        publication=pub,
        revision=revision,
        exclusions=exclusions,
        license_cell=license_cell,
        observations=observations,
    )
