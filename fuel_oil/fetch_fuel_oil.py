"""Reviewed replacement of the user-supplied candidate fetcher; offline by default.

Primary endpoints and legal interpretations remain unverified until access succeeds.
Raw snapshots are private, immutable, hash-bound and never silently admitted.
"""

import argparse
import calendar
import csv
from datetime import date, datetime, timedelta, timezone
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data/fuel_oil"
MAX_BYTES = 20 * 1024 * 1024
EIA_ROUTE = "https://api.eia.gov/v2/petroleum/pri/spt/data/"
FX_ROUTE = "https://data-api.ecb.europa.eu/service/data/EXR/"


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def encoded(value):
    return (
        json.dumps(value, indent=2, sort_keys=True, allow_nan=False) + "\n"
    ).encode()


def redact(value, secrets=()):
    text = str(value)
    for secret in secrets:
        if secret:
            text = text.replace(secret, "REDACTED").replace(
                urllib.parse.quote(secret, safe=""), "REDACTED"
            )
    return re.sub(
        r'(?i)((?:api[_-]?key|access_token|authorization)(?:=|%3d|:\s*))[^&\s"\'<>]+',
        r"\1REDACTED",
        text,
    )


def safe_url(url):
    u = urllib.parse.urlsplit(url)
    if u.scheme != "https" or not u.hostname or u.username or u.password or u.fragment:
        raise ValueError("HTTPS URL without credentials or fragment required")


def numeric(value, positive=False):
    if isinstance(value, bool) or value is None or str(value).strip() == "":
        raise ValueError("Missing numeric observation")
    try:
        v = float(value)
    except (TypeError, ValueError):
        raise ValueError("Invalid numeric observation") from None
    if not math.isfinite(v) or (positive and v <= 0):
        raise ValueError("Nonfinite or invalid positive observation")
    return v


def checked_period(value, frequency):
    patterns = {
        "monthly": r"[0-9]{4}-(0[1-9]|1[0-2])",
        "daily": r"[0-9]{4}-[0-9]{2}-[0-9]{2}",
    }
    if (
        frequency not in patterns
        or not isinstance(value, str)
        or not re.fullmatch(patterns[frequency], value)
    ):
        raise ValueError("Invalid observation period/frequency")
    date.fromisoformat(value + "-01" if frequency == "monthly" else value)
    return value


def unique(rows):
    keys = [r["period"] for r in rows]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate/revised period; select reviewed vintage upstream")
    if not rows:
        raise ValueError("No numeric observations")
    return sorted(rows, key=lambda r: r["period"])


def parse_fx(raw, frequency="monthly"):
    freq = {"monthly": "M", "daily": "D"}[frequency]
    expected = f"EXR.{freq}.USD.EUR.SP00.A"
    reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
    required = {
        "KEY",
        "FREQ",
        "CURRENCY",
        "CURRENCY_DENOM",
        "EXR_TYPE",
        "EXR_SUFFIX",
        "TIME_PERIOD",
        "OBS_VALUE",
        "OBS_STATUS",
        "UNIT",
        "UNIT_MULT",
    }
    if (
        not reader.fieldnames
        or not required.issubset(reader.fieldnames)
        or len(reader.fieldnames) != len(set(reader.fieldnames))
    ):
        raise ValueError("Unexpected FX layout")
    rows = []
    for r in reader:
        if None in r or None in r.values():
            raise ValueError("Malformed FX CSV")
        checks = {
            "KEY": expected,
            "FREQ": freq,
            "CURRENCY": "USD",
            "CURRENCY_DENOM": "EUR",
            "EXR_TYPE": "SP00",
            "EXR_SUFFIX": "A",
            "UNIT": "USD",
            "UNIT_MULT": "0",
        }
        if any(r[k] != v for k, v in checks.items()):
            raise ValueError(
                "FX key/unit/quote direction mismatch: requires USD per EUR"
            )
        if r["OBS_STATUS"] not in ("A", "P", "E"):
            raise ValueError("FX missing/unknown status; no fill")
        rows.append(
            dict(
                period=checked_period(r["TIME_PERIOD"], frequency),
                value=numeric(r["OBS_VALUE"], True),
                unit="USD_per_EUR",
                frequency=frequency,
                series_id=expected,
                source_status=r["OBS_STATUS"],
            )
        )
    return unique(rows)


def parse_eia_page(raw, frequency="monthly"):
    payload = json.loads(raw)
    if (
        not isinstance(payload, dict)
        or payload.get("error")
        or not isinstance(payload.get("response"), dict)
    ):
        raise ValueError("EIA error/unexpected payload")
    response = payload["response"]
    if response.get("frequency") != frequency:
        raise ValueError("EIA frequency mismatch")
    total = response.get("total")
    if isinstance(total, bool) or not re.fullmatch(r"[0-9]+", str(total)):
        raise ValueError("EIA total missing/invalid")
    data = response.get("data")
    if not isinstance(data, list) or not data:
        raise ValueError("EIA empty/non-numeric page")
    rows = []
    for r in data:
        if r.get("series") != "RBRTE" or r.get("units") != "Dollars per Barrel":
            raise ValueError("Wrong EIA series/units")
        rows.append(
            dict(
                period=checked_period(r.get("period"), frequency),
                value=numeric(r.get("value")),
                unit="USD_per_barrel",
                frequency=frequency,
                series_id="RBRTE",
            )
        )
    return int(total), unique(rows)


def combine_eia_pages(pages, offsets, frequency="monthly"):
    """Prove all declared result rows were retrieved; no implicit 5000-row ceiling."""
    rows, total = [], None
    if len(pages) != len(offsets) or not pages:
        raise ValueError("Missing pages/offsets")
    for raw, offset in zip(pages, offsets):
        count, batch = parse_eia_page(raw, frequency)
        if total is not None and count != total:
            raise ValueError("EIA total changed during paging; snapshot not coherent")
        if offset != len(rows):
            raise ValueError("EIA page coverage gap/overlap")
        total = count
        rows.extend(batch)
    if len(rows) != total:
        raise ValueError("EIA truncated/incomplete result")
    return unique(rows)


def parse_diesel_handoff(raw):
    """Explicit manually checked WOB mapping, NOT a verified native XLSX layout."""
    reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
    if reader.fieldnames != [
        "measurement_date",
        "country",
        "grade",
        "taxes",
        "unit",
        "value",
    ]:
        raise ValueError("Wrong diesel handoff layout")
    rows = []
    for r in reader:
        if None in r or None in r.values():
            raise ValueError("Malformed handoff")
        if (r["country"], r["grade"], r["taxes"], r["unit"]) != (
            "DE",
            "automotive_gas_oil",
            "with_taxes",
            "EUR_per_1000_litres",
        ):
            raise ValueError("Wrong country/grade/tax/unit; Euro-super 95 is never E10")
        p = checked_period(r["measurement_date"], "daily")
        if date.fromisoformat(p).weekday() != 0:
            raise ValueError(
                "Expected Monday measurement date; source convention must be verified"
            )
        value = numeric(r["value"], True)
        rows.append(
            dict(
                period=p,
                value=value,
                unit="EUR_per_1000_litres",
                frequency="weekly",
                country="DE",
                grade="automotive_gas_oil",
                taxes="with_taxes",
                eur_per_litre=value / 1000,
            )
        )
    return unique(rows)


def weekly_monthly(rows, start, end, as_of):
    """Monday value is an explicit approximation for the next seven calendar days.

    Not daily measurements, not interpolation over absent weeks, not official means.
    """
    unique(rows)
    covered = {}
    for r in rows:
        if (
            r.get("grade") != "automotive_gas_oil"
            or r.get("country") != "DE"
            or r.get("taxes") != "with_taxes"
            or r.get("unit") != "EUR_per_1000_litres"
        ):
            raise ValueError("Wrong diesel definition")
        d = date.fromisoformat(checked_period(r["period"], "daily"))
        if d.weekday() != 0:
            raise ValueError("Monday source measurement required")
        for k in range(7):
            day = d + timedelta(days=k)
            if day in covered:
                raise ValueError("Overlapping weekly intervals")
            covered[day] = numeric(r["value"], True) / 1000
    out = []
    for month in month_range(start, end):
        y, m = map(int, month.split("-"))
        days = [date(y, m, d) for d in range(1, calendar.monthrange(y, m)[1] + 1)]
        partial = days[-1] >= as_of
        present = [d for d in days if d in covered and d < as_of]
        valid = len(present) == len(days) and not partial
        out.append(
            dict(
                period=month,
                value=math.fsum(covered[d] for d in days) / len(days)
                if valid
                else None,
                unit="EUR_per_litre",
                covered_days=len(present),
                calendar_days=len(days),
                partial=partial,
                quality="derived_weekly_overlap_approximation"
                if valid
                else "incomplete",
                transformation="Monday-to-Sunday intervals weighted by overlap days; no absent-week fill",
            )
        )
    return out


def month_range(start, end):
    checked_period(start, "monthly")
    checked_period(end, "monthly")
    y, m = map(int, start.split("-"))
    out = []
    while f"{y:04d}-{m:02d}" <= end:
        out.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def monthly_coverage(rows, start, end, as_of):
    """Explicit gaps/partial periods, even when an API's total is fully paged."""
    unique(rows)
    if any(r.get("frequency") != "monthly" for r in rows):
        raise ValueError("Native monthly observations required")
    expected = month_range(start, end)
    observed = {checked_period(r["period"], "monthly") for r in rows}
    partial = sorted(p for p in observed if p >= as_of.strftime("%Y-%m"))
    return {
        "start": start,
        "end": end,
        "observed_months": len(set(expected) & observed),
        "missing_months": sorted(set(expected) - observed),
        "partial_or_future_months": partial,
        "historical_as_of": "not_established",
    }


def brent_eur_monthly(brent, fx, start, end, as_of, minimum_matched_days=15):
    """Predeclared >=15 dates AND complete overlap of the two supplied daily series.

    Calendar completeness of source exports needs separate source/paging evidence;
    matching alone cannot detect a day omitted by BOTH sources. No ratio of means.
    """
    unique(brent)
    unique(fx)
    if minimum_matched_days != 15:
        raise ValueError("Threshold fixed at 15; no data-driven choice")
    if any(
        r.get("frequency") != "daily" or r.get("unit") != "USD_per_barrel"
        for r in brent
    ) or any(
        r.get("frequency") != "daily" or r.get("unit") != "USD_per_EUR" for r in fx
    ):
        raise ValueError("Daily Brent and USD-per-EUR FX required")
    b = {checked_period(r["period"], "daily"): numeric(r["value"]) for r in brent}
    f = {checked_period(r["period"], "daily"): numeric(r["value"], True) for r in fx}
    out = []
    for month in month_range(start, end):
        bd = {d for d in b if d.startswith(month)}
        fd = {d for d in f if d.startswith(month)}
        matched = sorted(bd & fd)
        y, m = map(int, month.split("-"))
        partial = date(y, m, calendar.monthrange(y, m)[1]) >= as_of
        valid = len(matched) >= 15 and bd == fd and not partial
        out.append(
            dict(
                period=month,
                value=math.fsum(b[d] / f[d] for d in matched) / len(matched)
                if valid
                else None,
                unit="EUR_per_barrel",
                matched_days=len(matched),
                unmatched_brent_days=len(bd - fd),
                unmatched_fx_days=len(fd - bd),
                partial=partial,
                quality="derived_matched_daily_ratios"
                if valid
                else "insufficient_overlap_or_partial",
                transformation="mean(Brent_USD[d]/USD_per_EUR[d]); not ratio of monthly averages",
            )
        )
    return out


def check_xlsx(raw):
    """Container/layout quarantine only: no claimed WOB numeric observations."""
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if (
            sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024
            or "xl/workbook.xml" not in z.namelist()
        ):
            raise ValueError("Invalid/oversized XLSX container")
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        if not book.sheetnames:
            raise ValueError("Empty workbook")
    finally:
        book.close()
    return {
        "status": "quarantined_native_layout_not_verified",
        "numeric_observation_verified": False,
    }


def request(url, secrets=(), opener=None):
    """At most two attempts, transient HTTP 5xx only; no retry on 401/403/429."""
    safe_url(url)

    class NoRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *args):
            raise ValueError("Redirect requires reviewed endpoint")

    client = opener or urllib.request.build_opener(NoRedirect)
    record = dict(
        url=redact(url, secrets),
        requested_at=datetime.now(timezone.utc).isoformat(),
        retrieved_at=None,
        status=None,
        content_type=None,
        bytes=0,
        numeric_observation_verified=False,
        attempts=0,
    )
    for attempt in range(2):
        record["attempts"] = attempt + 1
        try:
            with client.open(
                urllib.request.Request(
                    url, headers={"User-Agent": "AAS-fuel-acquisition/1.0"}
                ),
                timeout=30,
            ) as r:
                record.update(
                    status=r.status,
                    content_type=r.headers.get("Content-Type", "")
                    .split(";")[0]
                    .strip()
                    .lower(),
                )
                if r.status != 200:
                    raise ValueError("Unexpected HTTP status")
                raw = r.read(MAX_BYTES + 1)
                if len(raw) > MAX_BYTES:
                    raise ValueError("Response size limit")
                # Never persist echoed secrets, even in a nominally valid response.
                if any(
                    s
                    and (
                        s.encode() in raw
                        or urllib.parse.quote(s, safe="").encode() in raw
                    )
                    for s in secrets
                ):
                    raise ValueError(
                        "Response echoed a credential; raw persistence refused"
                    )
                record.update(
                    bytes=len(raw),
                    raw_sha256=sha(raw),
                    retrieved_at=datetime.now(timezone.utc).isoformat(),
                )
                return raw, record
        except urllib.error.HTTPError as exc:
            record["status"] = exc.code
            if exc.code in (500, 502, 503, 504) and attempt == 0:
                time.sleep(0.25)
                continue
            record["error"] = redact(f"{type(exc).__name__}: {exc}", secrets)
            return None, record
        except Exception as exc:
            record["error"] = redact(f"{type(exc).__name__}: {exc}", secrets)
            return None, record
    raise AssertionError("Unreachable")


def validate_body(raw, content_type, kind, frequency):
    types = {
        "fx": {"text/csv", "application/vnd.sdmx.data+csv"},
        "eia": {"application/json"},
        "wob": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    }
    if content_type not in types[kind]:
        raise ValueError(
            "Unexpected content type; HTML/error response is not observations"
        )
    if kind == "fx":
        return parse_fx(raw, frequency)
    if kind == "eia":
        return parse_eia_page(raw, frequency)
    return check_xlsx(raw)


def rights_gate(review, fixture=False):
    fields = {
        "private_analysis",
        "raw_redistribution",
        "derived_redistribution",
        "review_url",
        "review_note",
        "attribution",
    }
    if set(review) != fields:
        raise ValueError("Incomplete rights review")
    safe_url(review["review_url"])
    if not review["review_note"].strip() or not review["attribution"].strip():
        raise ValueError("Rights evidence/attribution missing")
    expected = "synthetic_fixture" if fixture else "reviewed_permitted"
    if review["private_analysis"] != expected:
        raise ValueError("Private analysis not admitted")
    return review


def check_snapshot(raw, manifest, rights, fixture=False):
    if manifest.get("raw_sha256") != sha(raw):
        raise ValueError("Source hash mismatch")
    safe_url(manifest.get("url", ""))
    stamp = manifest.get("retrieved_at", "")
    if not re.fullmatch(
        r"[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}(?:\.[0-9]+)?(?:Z|\+00:00)",
        stamp,
    ):
        raise ValueError("Exact UTC retrieval timestamp required")
    datetime.fromisoformat(stamp.replace("Z", "+00:00"))
    if not manifest.get("vintage"):
        raise ValueError("Vintage required; historical availability remains separate")
    if bool(manifest.get("synthetic_fixture")) != fixture:
        raise ValueError("Synthetic/real boundary mismatch")
    rights_gate(rights, fixture)


def redistribution_gate(rights, derived=False):
    rights_gate(rights)
    field = "derived_redistribution" if derived else "raw_redistribution"
    if rights[field] != "reviewed_permitted":
        raise ValueError("Redistribution not reviewed/permitted; keep values private")


def save_raw(raw, record, out, name):
    if not re.fullmatch(r"[a-z0-9_]+\.(csv|json|xlsx)", name):
        raise ValueError("Contained snapshot filename required")
    path = Path(out) / name
    with path.open("xb") as f:
        f.write(raw)
    with path.with_suffix(path.suffix + ".manifest.json").open("xb") as f:
        f.write(
            encoded(
                dict(
                    record,
                    vintage="Retrospective retrieval; publisher revision unknown",
                    historical_as_of="not_established",
                    rights_status="not_verified",
                )
            )
        )


def fetch_eia(key, out, frequency="monthly", page_size=5000):
    if not key:
        raise ValueError("EIA_API_KEY not set; no fabricated key")
    pages, offsets, logs, offset, total = [], [], [], 0, None
    # Safety bound fails closed, never silently declares a truncated run successful.
    for page in range(100):
        params = {
            "frequency": frequency,
            "data[0]": "value",
            "facets[series][]": "RBRTE",
            "start": "2018-01-01" if frequency == "daily" else "2018-01",
            "sort[0][column]": "period",
            "sort[0][direction]": "asc",
            "offset": offset,
            "length": page_size,
            "api_key": key,
        }
        raw, rec = request(EIA_ROUTE + "?" + urllib.parse.urlencode(params), (key,))
        logs.append(rec)
        if raw is None:
            return None, logs
        try:
            count, rows = validate_body(raw, rec["content_type"], "eia", frequency)
            if total is not None and count != total:
                raise ValueError("EIA total changed")
            total = count
            rec.update(
                numeric_observation_verified=True,
                page_offset=offset,
                page_rows=len(rows),
                declared_total=total,
            )
            save_raw(raw, rec, out, f"eia_{frequency}_{page:03d}.json")
            pages.append(raw)
            offsets.append(offset)
            offset += len(rows)
            if offset >= total:
                all_rows = combine_eia_pages(pages, offsets, frequency)
                result = {
                    "observations": all_rows,
                    "declared_total": total,
                    "retrieved_rows": len(all_rows),
                    "page_offsets": offsets,
                }
                if frequency == "monthly":
                    result["coverage"] = monthly_coverage(
                        all_rows,
                        "2018-01",
                        "2025-12",
                        datetime.now(timezone.utc).date(),
                    )
                return result, logs
        except Exception as exc:
            rec["error"] = redact(
                f"Validation refused: {type(exc).__name__}: {exc}", (key,)
            )
            return None, logs
    logs[-1]["error"] = "Page safety limit; completeness not established"
    return None, logs


def blocked_coverage(sources):
    target = sources["target"]
    windows = [("2018-01", "2025-12"), ("2018-01", target["latest_completed_month"])]
    return {
        "sources_sha256": sha(encoded(sources)),
        "scope": "BLOCKED_ACQUISITION_NO_REAL_OBSERVATIONS",
        "empirical_completion": False,
        "series": [
            {
                "id": s["id"],
                "admitted": False,
                "reason": s["blockers"],
                "windows": [
                    {
                        "start": a,
                        "end": b,
                        "expected_months": len(month_range(a, b)),
                        "observed_months": 0,
                        "missing_or_unverified_months": len(month_range(a, b)),
                    }
                    for a, b in windows
                ],
            }
            for s in sources["series"]
        ],
        "partial_2026_months_excluded": target["partial_months_excluded"],
        "historical_as_of": "not_established",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--source", choices=("fx", "eia", "wob"))
    parser.add_argument("--frequency", choices=("daily", "monthly"), default="monthly")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    key = os.environ.get("EIA_API_KEY")
    try:
        sources = json.loads((DATA / "sources.json").read_text())
        if args.check:
            if (
                encoded(blocked_coverage(sources))
                != (DATA / "coverage.json").read_bytes()
            ):
                raise ValueError("Blocked coverage drift")
            print("Blocked acquisition verified; no real observations admitted")
            return
        if not args.source or not args.output:
            raise ValueError("Provide --source and a NEW private --output directory")
        args.output.mkdir(parents=True, exist_ok=False)
        if args.source == "eia":
            result, logs = fetch_eia(key, args.output, args.frequency)
        else:
            if args.source == "fx":
                freq = "D" if args.frequency == "daily" else "M"
                url = (
                    FX_ROUTE
                    + freq
                    + ".USD.EUR.SP00.A?startPeriod=2018-01&format=csvdata"
                )
            else:
                url = next(
                    s["download_candidate"]
                    for s in sources["series"]
                    if s["id"] == "diesel_de"
                )
            raw, rec = request(url, (key,) if key else ())
            logs, result = [rec], None
            if raw is not None:
                try:
                    parsed = validate_body(
                        raw, rec["content_type"], args.source, args.frequency
                    )
                    rec["numeric_observation_verified"] = args.source == "fx"
                    save_raw(
                        raw,
                        rec,
                        args.output,
                        "snapshot.csv" if args.source == "fx" else "snapshot.xlsx",
                    )
                    result = {
                        "verification": "numeric_snapshot_private_not_admitted"
                        if args.source == "fx"
                        else "quarantined_native_layout_not_verified",
                        "observation_count": len(parsed) if args.source == "fx" else 0,
                    }
                    if args.source == "fx" and args.frequency == "monthly":
                        result["coverage"] = monthly_coverage(
                            parsed,
                            "2018-01",
                            "2025-12",
                            datetime.now(timezone.utc).date(),
                        )
                except Exception as exc:
                    rec["error"] = redact(
                        f"Validation refused: {type(exc).__name__}: {exc}",
                        (key,) if key else (),
                    )
        (args.output / "access_log.json").write_bytes(encoded(logs))
        if result is None:
            raise ValueError("Acquisition incomplete; see redacted private access log")
        (args.output / "private_snapshot_status.json").write_bytes(
            encoded(
                {
                    "result": result,
                    "admitted": False,
                    "rights_status": "not_verified",
                    "historical_as_of": "not_established",
                }
            )
        )
        print(
            "Private snapshot only; no empirical admission or redistribution permission"
        )
    except Exception as exc:
        parser.exit(
            1,
            redact(
                f"Fuel/oil acquisition refused: {type(exc).__name__}: {exc}\n",
                (key,) if key else (),
            ),
        )


if __name__ == "__main__":
    main()
