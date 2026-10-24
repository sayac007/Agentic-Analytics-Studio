"""Actual user-export layouts: metadata-first, private quarantined QA, no models."""

import argparse
import calendar
import csv
from datetime import date, datetime, timezone, timedelta
import io
import json
from pathlib import Path
import re

import openpyxl
import xlrd

from fuel_oil.fetch_fuel_oil import (
    encoded,
    sha,
    numeric,
    month_range,
    unique,
    weekly_monthly,
)

ROOT = Path(__file__).resolve().parents[1]
KEY = "EXR.M.USD.EUR.SP00.A"
FX_HEADER = [
    "DATE",
    "TIME PERIOD",
    f"US dollar/Euro ECB reference exchange rate ({KEY})",
]
BRENT_TITLE = "Europe Brent Spot Price FOB (Dollars per Barrel)"
DIESEL_LABEL = "Gas oil automobile Automotive gas oil Dieselkraftstoff (I)"
NOTICE = "Reproduction is authorised provided the source is acknowledged."
MONTHS = {
    m: i
    for i, m in enumerate(
        (
            "Jan",
            "Feb",
            "Mar",
            "Apr",
            "May",
            "Jun",
            "Jul",
            "Aug",
            "Sep",
            "Oct",
            "Nov",
            "Dec",
        ),
        1,
    )
}


def value_status(value):
    if value in (None, "", "...", ":", ".."):
        return None, "missing_or_censored"
    return numeric(value, positive=True), "numeric_source_flags_not_provided"


def parse_fx(raw):
    reader = csv.reader(io.StringIO(raw.decode("utf-8-sig")))
    if next(reader, None) != FX_HEADER:
        raise ValueError("Wrong ECB portal key/currency/frequency/header")
    rows = []
    for line, r in enumerate(reader, 2):
        if len(r) != 3:
            raise ValueError("Malformed ECB portal row")
        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", r[0]):
            raise ValueError("Invalid ECB DATE")
        d = date.fromisoformat(r[0])
        expected = f"{d.year:04d}{tuple(MONTHS)[d.month - 1]}"
        if r[1] != expected or d.day != calendar.monthrange(d.year, d.month)[1]:
            raise ValueError("ECB date/month-label mismatch; month-end required")
        value, status = value_status(r[2])
        rows.append(
            dict(
                period=d.strftime("%Y-%m"),
                original_date=r[0],
                value=value,
                unit="USD_per_EUR",
                frequency="monthly",
                series_id=KEY,
                source_status=None,
                quality=status,
                source_location=f"CSV row {line}",
            )
        )
    return unique(rows)


def open_eia(raw):
    if not raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        raise ValueError("Expected native OLE XLS, not renamed XLSX/HTML")
    b = xlrd.open_workbook(file_contents=raw)
    if b.sheet_names() != ["Contents", "Data 1"]:
        raise ValueError("Unreviewed EIA workbook sheets")
    c, s = b.sheet_by_name("Contents"), b.sheet_by_name("Data 1")
    if (
        c.cell_value(2, 1) != BRENT_TITLE
        or c.cell_value(6, 4) != "Monthly"
        or c.cell_value(6, 2) != BRENT_TITLE
    ):
        raise ValueError("EIA identity/frequency mismatch")
    if (
        s.ncols != 2
        or s.cell_value(0, 1) != "Data 1: " + BRENT_TITLE
        or s.cell_value(1, 0) != "Sourcekey"
        or s.cell_value(1, 1) != "RBRTE"
        or s.row_values(2) != ["Date", BRENT_TITLE]
    ):
        raise ValueError("EIA series/units/layout mismatch")
    return b, c, s


def parse_brent(raw):
    b, _, s = open_eia(raw)
    rows = []
    for i in range(3, s.nrows):
        cell = s.cell(i, 0)
        if cell.ctype != xlrd.XL_CELL_DATE:
            raise ValueError(
                "EIA row is not a native monthly Excel date; footnote/annual summary refused"
            )
        d = xlrd.xldate.xldate_as_datetime(cell.value, b.datemode)
        value, status = value_status(s.cell_value(i, 1))
        rows.append(
            dict(
                period=d.strftime("%Y-%m"),
                original_date=d.date().isoformat(),
                value=value,
                frequency="monthly",
                unit="USD_per_barrel",
                series_id="RBRTE",
                source_status=None,
                quality=status,
                source_location=f"Data 1!A{i + 1}:B{i + 1}",
            )
        )
    return unique(rows)


def open_wob(raw):
    from fuel_oil.fetch_fuel_oil import check_xlsx

    check_xlsx(raw)
    b = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        if "Prices with taxes" not in b.sheetnames:
            raise ValueError("Missing with-tax sheet")
        s = b["Prices with taxes"]
        headers = list(s.iter_rows(min_row=1, max_row=3, values_only=True))
        if (
            headers[0][0]
            != "Consumer prices of petroleum products inclusive of duties and taxes"
            or headers[2][0] != "Date"
        ):
            raise ValueError("Wrong WOB tax/date definition")
        if sum(v == "DE_price_with_tax_diesel" for v in headers[0]) != 1:
            raise ValueError("Missing/duplicate Germany diesel key")
        # Pinned observed layout; moved fields require fresh source review.
        if (headers[0][55], headers[1][55], headers[2][55]) != (
            "DE_price_with_tax_diesel",
            DIESEL_LABEL,
            "1000 l",
        ):
            raise ValueError("Wrong WOB BD country/grade/tax/unit or moved layout")
        return b, s
    except Exception:
        b.close()
        raise


def parse_diesel(raw, unit_evidence=None):
    if unit_evidence is not None and (
        unit_evidence.get("unit") != "EUR_per_1000_litres"
        or unit_evidence.get("basis")
        != "user_attested_official_convention_and_independently_verified_mapping"
        or unit_evidence.get("citation")
        != "https://eur-lex.europa.eu/eli/dec/1999/566/oj/eng/pdf"
    ):
        raise ValueError("Unreviewed currency evidence")
    b, s = open_wob(raw)
    rows, notes, ended = [], [], False
    try:
        for line, r in enumerate(s.iter_rows(min_row=4, values_only=True), 4):
            d = r[0]
            if isinstance(d, datetime):
                if ended:
                    raise ValueError("Date rows after WOB footer/structure break")
                if (
                    r[53] != "DE_"
                    or d.time().isoformat() != "00:00:00"
                    or d.weekday() != 0
                ):
                    raise ValueError("WOB country/date convention drift")
                value, status = value_status(r[55])
                rows.append(
                    dict(
                        period=d.date().isoformat(),
                        original_date=d.date().isoformat(),
                        value=value,
                        frequency="weekly",
                        series_id="DE_price_with_tax_diesel",
                        unit="EUR_per_1000_litres"
                        if unit_evidence
                        else "currency_not_explicit_per_1000_litres",
                        country="DE",
                        grade="automotive_gas_oil",
                        taxes="with_taxes",
                        source_status=None,
                        quality=status,
                        source_location=f"Prices with taxes!A{line},BB{line},BD{line}",
                    )
                )
            elif d is not None and d != "Notes:":
                raise ValueError("Unexpected WOB date/structure")
            else:
                ended = True
                notes.extend(str(v) for v in r if isinstance(v, str))
        if not any(NOTICE in n for n in notes):
            raise ValueError("Expected embedded reproduction notice missing")
        if unit_evidence:
            for r in rows:
                r["eur_per_litre"] = None if r["value"] is None else r["value"] / 1000
                r["original_volume_unit"] = "1000 l"
        return unique(rows)
    finally:
        b.close()


def inspect_inputs(raws):
    """Inventory precedes extraction; no observed values in returned metadata."""
    fx, brent, diesel = (raws[k] for k in ("eurusd", "brent_usd", "diesel_de"))
    if next(csv.reader(io.StringIO(fx.decode("utf-8-sig"))), None) != FX_HEADER:
        raise ValueError("Unexpected ECB header")
    b, c, s = open_eia(brent)
    inventory = {
        "eurusd": {
            "header": FX_HEADER,
            "columns": 3,
            "data_rows": len(list(csv.reader(io.StringIO(fx.decode("utf-8-sig"))))) - 1,
        },
        "brent_usd": {
            "sheets": [
                {"name": x.name, "rows": x.nrows, "columns": x.ncols}
                for x in b.sheets()
            ],
            "header": BRENT_TITLE,
            "series_id": "RBRTE",
            "frequency": "Monthly",
            "publisher_workbook_timestamp_text": str(c.cell_value(14, 5)),
            "timestamp_meaning": "Workbook text only; timezone/publication/download not established",
        },
    }
    wb, sheet = open_wob(diesel)
    try:
        inventory["diesel_de"] = {
            "sheets": [
                {"name": x.title, "rows": x.max_row, "columns": x.max_column}
                for x in wb
            ],
            "selected_sheet": "Prices with taxes",
            "date_column": "A",
            "country_column": "BB",
            "price_column": "BD",
            "key": "DE_price_with_tax_diesel",
            "grade_label": DIESEL_LABEL,
            "volume_unit": "1000 l",
            "currency": "Not explicit in selected headers or cell format; EUR assumption not admitted",
            "embedded_rights_notice": NOTICE,
            "rights_locations": ["Prices with taxes!B1115", "Prices with taxes!B1117"],
        }
    finally:
        wb.close()
    return inventory


def coverage(rows, start, end, as_of):
    expected = month_range(start, end)
    weekly = rows[0]["frequency"] == "weekly"
    days = {}
    if weekly:
        for r in rows:
            if r["value"] is not None:
                d = date.fromisoformat(r["period"])
                for k in range(7):
                    p = d + timedelta(days=k)
                    if p in days:
                        raise ValueError("Weekly overlap")
                    days[p] = True
    by = {r["period"]: r for r in rows}
    observed = []
    missing = []
    censored = []
    partial = []
    for p in expected:
        y, m = map(int, p.split("-"))
        last = date(y, m, calendar.monthrange(y, m)[1])
        if last >= as_of:
            partial.append(p)
        if weekly:
            valid = all(date(y, m, d) in days for d in range(1, last.day + 1))
            if any(r["period"].startswith(p) and r["value"] is None for r in rows):
                censored.append(p)
        else:
            valid = p in by and by[p]["value"] is not None
            if p in by and by[p]["value"] is None:
                censored.append(p)
        (observed if valid and p not in partial else missing).append(p)
    return {
        "start": start,
        "end": end,
        "expected_months": len(expected),
        "observed_months": len(observed),
        "missing_months": len(missing),
        "missing_periods": missing,
        "censored_or_missing_value_months": len(censored),
        "partial_months": len(partial),
        "revised_months": None,
        "status_flags_unknown_months": len(expected),
        "admitted_months": 0,
        "blocked_months": len(expected),
        "measurement_basis": "Temporal seven-day interval coverage; monthly price admission is separate"
        if weekly
        else "Source-native monthly values",
    }


def run(paths, output, as_of, rights_log, fixture=False):
    output = Path(output).resolve()
    try:
        output.relative_to(ROOT.resolve())
    except ValueError:
        pass
    else:
        raise ValueError("Private outputs must be outside the Git repository")
    if output.exists():
        raise ValueError("Output exists; immutable runs cannot be overwritten")
    raws = {k: Path(v).read_bytes() for k, v in paths.items()}
    catalog_path = ROOT / (
        "tests/fixtures/energy_import/hashes.json"
        if fixture
        else "data/fuel_oil/real_sources.json"
    )
    catalog = json.loads(catalog_path.read_text())
    if catalog["scope"] != (
        "SYNTHETIC_FIXTURE_ONLY" if fixture else "USER_SUPPLIED_ORIGINALS"
    ) or set(raws) != set(catalog["files"]):
        raise ValueError("Source scope or inventory mismatch")
    for k, raw in raws.items():
        expected = catalog["files"][k]
        if (
            sha(raw) != expected["sha256"]
            or len(raw) != expected["bytes"]
            or Path(paths[k]).name != expected["filename"]
        ):
            raise ValueError("Original source filename/bytes/hash drift")
    inventory = inspect_inputs(raws)  # Must succeed BEFORE any numeric extraction.
    unit_evidence = catalog["files"]["diesel_de"].get("unit_evidence")
    parsers = {
        "eurusd": parse_fx,
        "brent_usd": parse_brent,
        "diesel_de": lambda raw: parse_diesel(raw, unit_evidence),
    }
    parsed = {k: fn(raws[k]) for k, fn in parsers.items()}
    summary = {
        "version": "1.0.0",
        "status": "SYNTHETIC_FIXTURE_ONLY"
        if fixture
        else "REAL_FILES_PARSED_NOT_ADMITTED",
        "source_catalog_sha256": sha(catalog_path.read_bytes()),
        "as_of": as_of.isoformat(),
        "historical_as_of": "not_established",
        "source_acquisition": "User-supplied originals; not successful live API retrieval",
        "rights_review_requests": rights_log,
        "series": [],
    }
    for k, rows in parsed.items():
        for r in rows:
            r["partial_calendar_month"] = r["period"][:7] >= as_of.strftime("%Y-%m")
        inventory[k].update(
            filename=Path(paths[k]).name,
            bytes=len(raws[k]),
            raw_sha256=sha(raws[k]),
            actual_download_time=None,
            local_ingestion_is_not_download=True,
        )
        gaps = [
            "Actual source download URL and retrieval evidence not provided",
            "Source status/revision flags not provided",
        ]
        reasons = []
        if k == "diesel_de":
            if unit_evidence:
                inventory[k]["currency"] = (
                    "EUR per 1000 litres accepted from user-cited official convention plus independently checked mapping; primary PDF fetch proxy-blocked"
                )
                inventory[k]["unit_evidence"] = unit_evidence
                reasons.append(
                    "Derived-use scope not independently verified; private QA normalization is not final analytical admission"
                )
            else:
                reasons.append(
                    "Currency not explicit in selected workbook headers; EUR conversion and price aggregation blocked"
                )
        else:
            reasons.append(
                "Primary rights notice unreachable; private/raw/derived permission not verified"
            )
        summary["series"].append(
            {
                "id": k,
                "inventory": inventory[k],
                "admitted": False,
                "rights": {
                    "private_analysis": "reviewed_embedded_reproduction_notice"
                    if k == "diesel_de"
                    else "not_verified",
                    "raw_redistribution": "conditional_embedded_notice_keep_private"
                    if k == "diesel_de"
                    else "not_verified",
                    "derived_redistribution": "not_verified",
                },
                "blockers": reasons,
                "provenance_gaps": gaps,
                "native_rows": len(rows),
                "native_start": rows[0]["period"],
                "native_end": rows[-1]["period"],
                "windows": [
                    coverage(rows, a, b, as_of)
                    for a, b in [("2018-01", "2025-12"), ("2018-01", "2026-09")]
                ],
                "partial_october_2026_native_rows": sum(
                    r["period"].startswith("2026-10") for r in rows
                ),
                "duplicate_periods": 0,
                "nonfinite_values": 0,
            }
        )
    output.mkdir(parents=True)
    artifacts = {}
    for k in parsed:
        raw_name = Path(paths[k]).name
        path = output / (k + Path(raw_name).suffix)
        with path.open("xb") as f:
            f.write(raws[k])
        payload = {
            "status": "QUARANTINED_STRUCTURAL_QA_ONLY_NOT_ADMITTED",
            "raw_sha256": sha(raws[k]),
            "actual_download_time": None,
            "source_vintage": "Exact uploaded bytes; publisher revision not established",
            "historical_as_of": "not_established",
            "observations": parsed[k],
        }
        normalized = encoded(payload)
        with (output / (k + ".quarantined.json")).open("xb") as f:
            f.write(normalized)
        artifacts[k] = {
            "raw_sha256": sha(raws[k]),
            "normalized_sha256": sha(normalized),
        }
    if unit_evidence:
        monthly = weekly_monthly(parsed["diesel_de"], "2018-01", "2026-10", as_of)
        payload = encoded(
            {
                "status": "QUARANTINED_PRIVATE_DERIVED_QA_NOT_ADMITTED",
                "raw_sha256": sha(raws["diesel_de"]),
                "unit_evidence": unit_evidence,
                "observations": monthly,
                "warning": "Derived seven-day overlap approximation, not official monthly means; missing/partial months null",
            }
        )
        (output / "diesel_de.monthly_approx.quarantined.json").write_bytes(payload)
        artifacts["diesel_monthly_approx"] = {
            "raw_sha256": sha(raws["diesel_de"]),
            "normalized_sha256": sha(payload),
        }
    (output / "inventory.json").write_bytes(
        encoded({"artifacts": artifacts, "summary_sha256": sha(encoded(summary))})
    )
    (output / "summary.json").write_bytes(encoded(summary))
    (output / "local_ingestion.json").write_bytes(
        encoded(
            {
                "ingested_at": datetime.now(timezone.utc).isoformat(),
                "meaning": "Own ingestion only, not user/provider download time",
            }
        )
    )
    return summary


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--ecb", type=Path)
    p.add_argument("--brent", type=Path)
    p.add_argument("--diesel", type=Path)
    p.add_argument("--output", type=Path)
    p.add_argument("--as-of", type=date.fromisoformat)
    p.add_argument("--rights-log", type=Path)
    p.add_argument("--fixture", action="store_true")
    p.add_argument("--check-metadata", action="store_true")
    a = p.parse_args()
    try:
        if a.check_metadata:
            check_public_metadata()
            print("Metadata-only real upload summary verified; no private files read")
            return
        if not all((a.ecb, a.brent, a.diesel, a.output, a.as_of, a.rights_log)):
            raise ValueError(
                "Source files, private output, as-of and rights log required"
            )
        run(
            {"eurusd": a.ecb, "brent_usd": a.brent, "diesel_de": a.diesel},
            a.output,
            a.as_of,
            json.loads(a.rights_log.read_text()),
            fixture=a.fixture,
        )
        print("Private structural QA complete; all empirical admission remains blocked")
    except Exception as e:
        # Do not expose user local paths in failure messages.
        p.exit(
            1,
            f"Energy import refused ({type(e).__name__}); no public values produced\n",
        )


def check_public_metadata():
    catalog_path = ROOT / "data/fuel_oil/real_sources.json"
    summary = json.loads((ROOT / "data/fuel_oil/real_coverage.json").read_text())
    catalog = json.loads(catalog_path.read_text())
    if summary["source_catalog_sha256"] != sha(catalog_path.read_bytes()):
        raise ValueError("Source-catalog binding drift")

    def safe(value):
        if isinstance(value, dict):
            if {"observations", "value", "local_path", "private_path"} & set(value):
                raise ValueError(
                    "Private observation/path fields forbidden in public summary"
                )
            for x in value.values():
                safe(x)
        elif isinstance(value, list):
            for x in value:
                safe(x)

    safe(summary)
    for s in summary["series"]:
        record = catalog["files"][s["id"]]
        if s["admitted"] or s["inventory"]["raw_sha256"] != record["sha256"]:
            raise ValueError("Unreviewed admission or raw hash drift")
        for w in s["windows"]:
            if (
                w["expected_months"] != len(month_range(w["start"], w["end"]))
                or w["observed_months"] + w["missing_months"] != w["expected_months"]
                or w["admitted_months"] != 0
            ):
                raise ValueError("Invalid metadata coverage counts")


if __name__ == "__main__":
    main()
